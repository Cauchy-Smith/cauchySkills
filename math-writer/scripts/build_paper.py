#!/usr/bin/env python3
"""Build the competition paper .docx from a Markdown source.

The builder owns typography; the source owns content. Formatting follows
《"华为杯"第二十三届中国研究生数学建模竞赛论文格式规范》:

    论文题目  三号黑体, 居中
    一级标题  四号黑体, 居中
    其他汉字  小四号宋体
    行距      单倍行距
    页脚      居中阿拉伯数字, 从摘要页起为 1, 封面不编号
    页眉      无

Equations become real OMML (never images) via latex2omml.

Usage:
    python build_paper.py paper/paper.md -o paper/论文.docx

See references/source-format.md for the Markdown subset this accepts.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from latex2omml import LatexError, to_omml  # noqa: E402

# ---- typography constants (Chinese point sizes) --------------------------- #
SIZE_TITLE = Pt(16)      # 三号
SIZE_H1 = Pt(14)         # 四号
SIZE_BODY = Pt(12)       # 小四
SIZE_SMALL = Pt(10.5)    # 五号
SIZE_COVER = Pt(18)      # 小二

FONT_HEI = "黑体"
FONT_SONG = "宋体"
FONT_EN = "Times New Roman"
FONT_CODE = "Consolas"

PAGE_W_CM, PAGE_H_CM = 21.0, 29.7            # A4
MARGIN_L_CM = MARGIN_R_CM = 2.25
MARGIN_T_CM, MARGIN_B_CM = 3.0, 1.75
USABLE_W_CM = PAGE_W_CM - MARGIN_L_CM - MARGIN_R_CM

PAGE_W, PAGE_H = Cm(PAGE_W_CM), Cm(PAGE_H_CM)
MARGIN_L = MARGIN_R = Cm(MARGIN_L_CM)
MARGIN_T, MARGIN_B = Cm(MARGIN_T_CM), Cm(MARGIN_B_CM)

UNNUMBERED_HEADINGS = {"参考文献", "附录", "致谢", "摘要"}

# "表：题注" 或 "表2-1：题注"（后者使用显式编号）
TABLE_CAP_RE = re.compile(r"^表\s*(?:\d+\s*[-–]\s*\d+)?\s*[:：]")


# --------------------------------------------------------------------------- #
# low-level docx helpers
# --------------------------------------------------------------------------- #

def style_run(run, *, cn=FONT_SONG, en=FONT_EN, size=SIZE_BODY, bold=False,
              italic=False, color=None):
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)
    if color is not None:
        run.font.color.rgb = color
    return run


def set_spacing(paragraph, *, first_line_chars=None, line=WD_LINE_SPACING.SINGLE,
                before=0, after=0, keep_with_next=False):
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if keep_with_next:
        pf.keep_with_next = True
    if first_line_chars:
        ppr = paragraph._p.get_or_add_pPr()
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            ppr.append(ind)
        ind.set(qn("w:firstLineChars"), str(first_line_chars))
        ind.set(qn("w:firstLine"), str(int(first_line_chars * 2.4)))  # fallback twips-ish


def add_field(paragraph, instr: str, placeholder: str = "1"):
    """Insert a Word field (e.g. PAGE) into a paragraph."""
    for kind, payload in (("begin", None), (None, instr), ("separate", None),
                          (None, placeholder), ("end", None)):
        run = paragraph.add_run()
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        else:
            if payload == instr:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = instr
            else:
                el = OxmlElement("w:t")
                el.text = payload
            run._r.append(el)


def restart_page_numbering(section, start: int = 1):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectPr.append(pg)
    pg.set(qn("w:start"), str(start))


def set_cell_text(cell, text, *, bold=False, size=SIZE_BODY, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    set_spacing(p)
    style_run(p.add_run(text), cn=FONT_SONG, size=size, bold=bold)


def set_cell_inline(cell, text, *, bold=False, size=SIZE_BODY, align=WD_ALIGN_PARAGRAPH.CENTER):
    """Same as set_cell_text but renders inline $math$ / **bold** / citations."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    set_spacing(p)
    add_inline(p, text, size=size, bold=bold)


def three_line_borders(table):
    """Top rule + header rule + bottom rule, nothing else (三线表)."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, val, sz in (("top", "single", 12), ("bottom", "single", 12),
                          ("left", "none", 0), ("right", "none", 0),
                          ("insideH", "none", 0), ("insideV", "none", 0)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)

    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "0")
        bottom.set(qn("w:color"), "000000")
        tc_borders.append(bottom)
        tc_pr.append(tc_borders)


# --------------------------------------------------------------------------- #
# inline parsing: $math$, **bold**, ^[1] citations
# --------------------------------------------------------------------------- #

INLINE_RE = re.compile(
    r"(\$[^$]+\$|\*\*[^*]+\*\*|`[^`]+`|\^\[[0-9][0-9,\-\u2013\s]*\])")


def add_inline(paragraph, text: str, *, size=SIZE_BODY, bold=False):
    """Append mixed text/math/bold/inline-code/citation content to a paragraph."""
    from lxml import etree

    for piece in INLINE_RE.split(text):
        if not piece:
            continue
        if piece.startswith("$") and piece.endswith("$") and len(piece) > 2:
            try:
                paragraph._p.append(etree.fromstring(to_omml(piece[1:-1])))
            except LatexError:
                style_run(paragraph.add_run(piece), size=size, bold=bold)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            style_run(paragraph.add_run(piece[1:-1]),
                      cn=FONT_SONG, en=FONT_CODE, size=SIZE_SMALL, bold=bold)
        elif piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            add_inline(paragraph, piece[2:-2], size=size, bold=True)
        elif piece.startswith("^[") and piece.endswith("]"):
            run = style_run(paragraph.add_run(piece[1:]), size=size, bold=bold)
            run.font.superscript = True
        else:
            style_run(paragraph.add_run(piece), size=size, bold=bold)


# --------------------------------------------------------------------------- #
# document setup
# --------------------------------------------------------------------------- #

def page_setup(section):
    section.page_width = PAGE_W
    section.page_height = PAGE_H
    section.left_margin = MARGIN_L
    section.right_margin = MARGIN_R
    section.top_margin = MARGIN_T
    section.bottom_margin = MARGIN_B
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.0)


def make_document(template=None):
    if template:
        doc = docx.Document(str(template))
        body = doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):     # keep page geometry/styles
                body.remove(child)
    else:
        doc = docx.Document()
        page_setup(doc.sections[0])
    normal = doc.styles["Normal"]
    normal.font.name = FONT_EN
    normal.font.size = SIZE_BODY
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_SONG)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.space_after = Pt(0)
    return doc


# --------------------------------------------------------------------------- #
# sections of the paper
# --------------------------------------------------------------------------- #

def build_cover(doc, meta):
    for _ in range(2):
        doc.add_paragraph()
    banner = meta.get("banner")
    if banner:
        lines = [ln.strip() for ln in banner.split("|") if ln.strip()]
    else:
        lines = [meta.get("competition", "中国研究生创新实践系列大赛"),
                 meta.get("edition", "“华为杯”第二十三届中国研究生"),
                 meta.get("contest_name", "数学建模竞赛")]
    for text in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p, before=6, after=6)
        style_run(p.add_run(text), cn=FONT_HEI, size=SIZE_COVER, bold=True)

    for _ in range(3):
        doc.add_paragraph()

    rows = [("学　　校", meta.get("school", "")),
            ("参赛队号", meta.get("team_id", "")),
            ("队员姓名", meta.get("members", ""))]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, value) in enumerate(rows):
        set_cell_text(table.cell(i, 0), label, size=SIZE_BODY)
        set_cell_text(table.cell(i, 1), value, size=SIZE_BODY)


def insert_toc(doc, levels: str = "1-3"):
    """Insert a Word TOC field. Word fills it in when fields are updated
    (export_pdf.ps1 does that), so the entry list always matches the headings.
    """
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(title, before=0, after=12)
    style_run(title.add_run("目　录"), cn=FONT_HEI, size=SIZE_TITLE, bold=True,
              color=RGBColor(0, 0, 0))

    para = doc.add_paragraph()
    set_spacing(para, after=0)
    for kind, payload in (("begin", None),
                          (None, f' TOC \\o "{levels}" \\h \\z \\u '),
                          ("separate", None),
                          (None, "（在 Word 中按 F9 或导出 PDF 时自动生成目录）"),
                          ("end", None)):
        run = para.add_run()
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                fld.set(qn("w:dirty"), "true")
            run._r.append(fld)
        else:
            el = OxmlElement("w:instrText" if payload.startswith(" TOC") else "w:t")
            el.set(qn("xml:space"), "preserve")
            el.text = payload
            run._r.append(el)
            if not payload.startswith(" TOC"):
                style_run(run, size=SIZE_SMALL)


def build_abstract_page(doc, meta, abstract_paras):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, before=0, after=12)
    style_run(p.add_run(f"题目：{meta.get('title', '')}"),
              cn=FONT_HEI, size=SIZE_TITLE, bold=True)

    label = doc.add_paragraph()
    set_spacing(label, before=6, after=6)
    style_run(label.add_run("摘　要："), cn=FONT_HEI, size=SIZE_H1, bold=True)

    for text in abstract_paras:
        para = doc.add_paragraph()
        set_spacing(para, first_line_chars=200)
        add_inline(para, text)

    kw = doc.add_paragraph()
    set_spacing(kw, before=12)
    style_run(kw.add_run("关键词："), bold=True)
    style_run(kw.add_run(meta.get("keywords", "")))


# --------------------------------------------------------------------------- #
# markdown body
# --------------------------------------------------------------------------- #

class BodyBuilder:
    def __init__(self, doc, base: Path | None = None):
        self.doc = doc
        self.base = Path(base) if base else Path.cwd()
        self.chapter = 0
        self.section = 0
        self.subsection = 0
        self.fig_chapter = None
        self.fig_n = 0
        self.tbl_n = 0
        self.eq_n = 0
        self.pending_table_caption = None

    # -- counters ---------------------------------------------------------- #
    def _reset_counters(self):
        self.section = 0
        self.subsection = 0
        self.fig_n = 0
        self.tbl_n = 0
        self.eq_n = 0
        self.fig_chapter = self.chapter

    def next_fig(self):
        self._ensure_chapter()
        self.fig_n += 1
        return f"图{self.fig_chapter}-{self.fig_n}"

    def next_tbl(self):
        self._ensure_chapter()
        self.tbl_n += 1
        return f"表{self.fig_chapter}-{self.tbl_n}"

    def next_eq(self):
        self._ensure_chapter()
        self.eq_n += 1
        return f"({self.fig_chapter}-{self.eq_n})"

    def _ensure_chapter(self):
        if self.fig_chapter != self.chapter:
            self.fig_chapter = self.chapter
            self.fig_n = self.tbl_n = self.eq_n = 0

    # -- blocks ------------------------------------------------------------ #
    def heading(self, level, text, numbered=True):
        if level == 1:
            self.chapter += 1
            self._reset_counters()
            prefix = self._number_prefix(text, [self.chapter]) if numbered else ""
        elif level == 2:
            self.section += 1
            self.subsection = 0
            prefix = self._number_prefix(text, [self.chapter, self.section]) if numbered else ""
        else:
            self.subsection += 1
            prefix = self._number_prefix(
                text, [self.chapter, self.section, self.subsection]) if numbered else ""

        p = self.doc.add_paragraph()
        # Heading styles carry the outline level; the TOC field needs them.
        # Run-level formatting below overrides the style's own font/colour.
        # Templates without named heading styles get an explicit outline level.
        style = None
        for name in (f"Heading {min(level, 3)}", f"标题 {min(level, 3)}"):
            try:
                style = self.doc.styles[name]
                break
            except KeyError:
                continue
        if style is not None:
            p.style = style
        else:
            lvl = OxmlElement("w:outlineLvl")
            lvl.set(qn("w:val"), str(min(level, 3) - 1))
            p._p.get_or_add_pPr().append(lvl)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        set_spacing(p, before=12 if level == 1 else 8, after=6)
        size = SIZE_H1 if level == 1 else SIZE_BODY
        cn = FONT_HEI if level == 1 else FONT_SONG
        style_run(p.add_run(f"{prefix}{text}".strip()), cn=cn, size=size, bold=True,
                  color=RGBColor(0, 0, 0))
        return p

    @staticmethod
    def _number_prefix(text, nums):
        """Skip auto numbering when the author already wrote one."""
        if re.match(r"^\s*\d+(\.\d+)*[\s、.]", text) or re.match(r"^\s*[一二三四五六七八九十]+[、.]", text):
            return ""
        return ".".join(str(n) for n in nums) + " "

    def paragraph(self, text):
        p = self.doc.add_paragraph()
        set_spacing(p, first_line_chars=200)
        add_inline(p, text)

    def bullet(self, text):
        p = self.doc.add_paragraph(style="List Bullet")
        set_spacing(p)
        add_inline(p, text)

    def numbered(self, text):
        p = self.doc.add_paragraph(style="List Number")
        set_spacing(p)
        add_inline(p, text)

    def display_math(self, latex):
        """Centered equation with a right-aligned number, via tab stops."""
        import re as _re
        m = _re.search(r"\\tag\{([^}]*)\}", latex)
        if m:
            raw = m.group(1).strip()
            # emitted numbers are parenthesised, so a bare \tag{5-1} must not
            # render as "5-1" while auto numbers render as "(5-1)"
            number = raw if raw.startswith("(") else f"({raw})"
            latex = latex[:m.start()] + latex[m.end():]
        else:
            number = self.next_eq()
        p = self.doc.add_paragraph()
        set_spacing(p, before=6, after=6)
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(USABLE_W_CM / 2), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(USABLE_W_CM), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run("\t")
        from lxml import etree
        p._p.append(etree.fromstring(to_omml(latex)))
        p.add_run("\t")
        style_run(p.add_run(number), size=SIZE_BODY)

    def figure(self, path, caption):
        # an explicit leading 图<ch>-<n> wins over auto-numbering, so a figure
        # can follow the coder's per-question filename instead of the chapter
        label, caption = self._explicit_number("图", caption) or (self.next_fig(), caption)
        candidate = Path(path)
        if not candidate.is_absolute():
            candidate = (self.base / candidate).resolve()
        path = str(candidate)
        pic_para = self.doc.add_paragraph()
        pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(pic_para, before=6, after=3, keep_with_next=True)
        if Path(path).exists():
            pic_para.add_run().add_picture(path, width=Cm(12))
        else:
            style_run(pic_para.add_run(f"[缺失图片: {path}]"), size=SIZE_SMALL,
                      color=RGBColor(0xC0, 0x00, 0x00))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(cap, after=10)
        style_run(cap.add_run(f"{label}　{caption}".rstrip("　")), size=SIZE_SMALL)

    @staticmethod
    def _explicit_number(kind, caption):
        m = re.match(rf"^\s*{kind}\s*(\d+)\s*[-–]\s*(\d+)\s*[:：]?\s*(.*)$", caption)
        if not m:
            return None
        return f"{kind}{m.group(1)}-{m.group(2)}", m.group(3)

    def table(self, rows, caption=None):
        cap_text = caption or self.pending_table_caption
        self.pending_table_caption = None
        cap_text = cap_text or ""
        explicit = self._explicit_number("表", cap_text)
        label, cap_text = explicit or (self.next_tbl(), cap_text)
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(cap, before=8, after=3, keep_with_next=True)
        style_run(cap.add_run(f"{label}　{cap_text}".rstrip("　")), size=SIZE_SMALL)

        n_cols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=n_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for i, row in enumerate(rows):
            for j in range(n_cols):
                text = row[j] if j < len(row) else ""
                set_cell_inline(table.cell(i, j), text, bold=(i == 0), size=SIZE_SMALL)
        three_line_borders(table)

    def code_block(self, lines):
        for line in lines:
            p = self.doc.add_paragraph()
            set_spacing(p)
            style_run(p.add_run(line), cn=FONT_SONG, en=FONT_CODE, size=SIZE_SMALL)

    def reference(self, text):
        p = self.doc.add_paragraph()
        set_spacing(p)
        pf = p.paragraph_format
        pf.left_indent = Cm(0.85)
        pf.first_line_indent = Cm(-0.85)
        add_inline(p, text, size=SIZE_SMALL)


# --------------------------------------------------------------------------- #
# source parsing
# --------------------------------------------------------------------------- #

def parse_front_matter(source: str):
    meta, body = {}, source
    if source.lstrip().startswith("---"):
        parts = source.split("---", 2)
        if len(parts) >= 3:
            block, body = parts[1], parts[2]
            for line in block.splitlines():
                if ":" in line:
                    k, v = line.split(":", 1)
                    meta[k.strip()] = v.strip().strip('"').strip("'")
    return meta, body.lstrip("\n")


def parse_pipe_row(line: str):
    """Split a Markdown pipe row, keeping pipes inside $...$ and honoring \\|."""
    s = line.strip().strip("|")
    cells, buf, in_math, i = [], [], False, 0
    while i < len(s):
        ch = s[i]
        if ch == "\\" and i + 1 < len(s) and s[i + 1] == "|":
            buf.append("\\|" if in_math else "|")
            i += 2
            continue
        if ch == "$":
            in_math = not in_math
        if ch == "|" and not in_math:
            cells.append("".join(buf).strip())
            buf = []
        else:
            buf.append(ch)
        i += 1
    cells.append("".join(buf).strip())
    return cells


def is_separator_row(cells):
    return all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c != "")


def build_body(doc, body_md: str, base=None):
    b = BodyBuilder(doc, base=base)
    lines = body_md.splitlines()
    i = 0
    in_refs = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            b.code_block(block)
            continue

        if stripped.startswith("$$"):
            rest = stripped[2:]
            if "$$" in rest:                      # single-line $$ ... $$
                latex = rest.split("$$")[0]
                i += 1
            else:                                  # fenced display block
                parts = [rest] if rest else []
                i += 1
                while i < len(lines):
                    if "$$" in lines[i]:
                        parts.append(lines[i].split("$$")[0])
                        i += 1
                        break
                    parts.append(lines[i].strip())
                    i += 1
                latex = " ".join(p for p in parts if p)
            if not latex.strip():
                print("[WARN] empty display equation skipped", file=sys.stderr)
                continue
            b.display_math(latex.strip())
            continue

        if stripped.startswith("!["):
            m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
            if m:
                b.figure(m.group(2).strip(), m.group(1).strip())
                i += 1
                continue

        if TABLE_CAP_RE.match(stripped) or stripped.startswith(("表：", "表:")):
            b.pending_table_caption = stripped[2:].strip() if stripped[1] in "：:" else stripped
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = parse_pipe_row(lines[i])
                if not is_separator_row(cells):
                    rows.append(cells)
                i += 1
            if rows:
                b.table(rows)
            continue

        if re.match(r"^#{1,3}\s", stripped):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if text in ("摘要", "关键词"):
                i += 1
                continue
            numbered = text not in UNNUMBERED_HEADINGS
            b.heading(level, text, numbered=numbered)
            in_refs = text == "参考文献"
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            b.bullet(re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue

        if re.match(r"^\d+[.)]\s+", stripped) and not in_refs:
            b.numbered(re.sub(r"^\d+[.)]\s+", "", stripped))
            i += 1
            continue

        if in_refs and re.match(r"^\[\d+\]", stripped):
            b.reference(stripped)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        para = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", "```", "![", "$$"))
                    or TABLE_CAP_RE.match(nxt)
                    or re.match(r"^[-*]\s+", nxt)):
                break
            if in_refs and re.match(r"^\[\d+\]", nxt):
                break
            para.append(nxt)
            i += 1
        b.paragraph(" ".join(para))


def split_abstract(body_md: str):
    """Return (abstract paragraphs, keywords found in the abstract, remaining body).

    The abstract belongs on the abstract page only; leaving it in the body
    would print it twice. A `关键词：...` line inside the 摘要 section is
    pulled out too, otherwise build_abstract_page would print it a second time.
    """
    lines = body_md.splitlines()
    abstract, keywords, kept, i = [], None, [], 0
    while i < len(lines):
        if re.match(r"^#\s*摘要\s*$", lines[i].strip()):
            i += 1
            while i < len(lines) and not re.match(r"^#{1,3}\s", lines[i].strip()):
                s = lines[i].strip()
                if s:
                    m = re.match(r"^关键词\s*[:：]\s*(.*)$", s)
                    if m:
                        keywords = m.group(1).strip()
                    else:
                        abstract.append(s)
                i += 1
            continue
        kept.append(lines[i])
        i += 1
    return abstract, keywords, "\n".join(kept)


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #

def build(source: Path, out: Path, template: Path | None = None):
    text = source.read_text(encoding="utf-8")
    meta, body_md = parse_front_matter(text)
    abstract_paras, body_keywords, body_md = split_abstract(body_md)
    abstract = meta.pop("abstract", None)
    if abstract:
        abstract_paras = [abstract]
    if not meta.get("keywords") and body_keywords:
        meta["keywords"] = body_keywords

    doc = make_document(template)
    build_cover(doc, meta)

    cover = doc.sections[0]
    cover.footer.is_linked_to_previous = False
    cover.footer.paragraphs[0].text = ""

    body_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(body_sec)
    body_sec.header.is_linked_to_previous = False
    body_sec.header.paragraphs[0].text = ""
    body_sec.footer.is_linked_to_previous = False
    restart_page_numbering(body_sec, 1)
    footer_p = body_sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.text = ""
    add_field(footer_p, " PAGE ", "1")

    build_abstract_page(doc, meta, abstract_paras)
    doc.add_page_break()

    want_toc = str(meta.get("toc", "true")).strip().lower() not in ("false", "no", "0")
    if want_toc:
        insert_toc(doc, str(meta.get("toc_levels", "1-3")))
        meta["_has_toc"] = True
        doc.add_page_break()

    build_body(doc, body_md, base=source.parent)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    write_build_manifest(source, out, template)
    return doc


def write_build_manifest(source: Path, out: Path, template=None):
    """Bind the built .docx to its source, template and assets so a stale
    artefact can be detected (any input change invalidates the build)."""
    import datetime
    import hashlib
    import json

    def sha(path):
        p = Path(path)
        return hashlib.sha256(p.read_bytes()).hexdigest() if p.exists() else None

    assets = {}
    for path in sorted(source.parent.rglob("*")):
        if path.is_file() and path.resolve() != Path(out).resolve():
            assets[str(path.relative_to(source.parent)).replace("\\", "/")] = sha(path)

    manifest = {
        "generated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "source": {"path": str(source), "sha256": sha(source)},
        "template": ({"path": str(template), "sha256": sha(template)}
                     if template else None),
        "output": {"path": str(out), "sha256": sha(out)},
        "assets": assets,
    }
    Path(f"{out}.build.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("source", help="Markdown source")
    ap.add_argument("-o", "--out", required=True, help="Output .docx path")
    ap.add_argument("--template", help="Official .docx template to inherit styles/sections from")
    args = ap.parse_args(argv)

    src = Path(args.source)
    if not src.exists():
        print(f"[ERROR] source not found: {src}", file=sys.stderr)
        return 1
    template = Path(args.template) if args.template else None
    if template and not template.exists():
        print(f"[ERROR] template not found: {template}", file=sys.stderr)
        return 1
    try:
        build(src, Path(args.out), template)
    except LatexError as exc:
        print(f"[ERROR] equation failed to convert: {exc}", file=sys.stderr)
        return 2
    print(f"[ok] wrote {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
